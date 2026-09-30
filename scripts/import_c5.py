"""ARCHIVED: former PostgreSQL/WiFi C5 import, retained only as reference.

Not part of the active C5 flow. Its CLI is disabled to prevent accidental imports.
"""
import argparse
import csv
import json
import math
import sys
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from sqlalchemy import text
from redis.exceptions import RedisError
from app.core.database import get_engine
from app.services.c5_repository import rebuild_cache, IMPORT_LOCK


def normalized(value):
    return "_".join("".join(c for c in unicodedata.normalize("NFKD", str(value).strip().lower()) if not unicodedata.combining(c)).split())


def rows(path, encoding):
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = book.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values)
            for row in values:
                yield dict(zip(headers, row))
        finally:
            book.close()
    elif path.suffix.lower() in {".json", ".geojson"}:
        data = json.loads(path.read_text(encoding=encoding))
        if data.get("type") != "FeatureCollection":
            raise ValueError("Se requiere GeoJSON FeatureCollection")
        for feature in data["features"]:
            row = dict(feature.get("properties") or {})
            geometry = feature.get("geometry") or {}
            if geometry.get("type") == "Point":
                row["longitude"], row["latitude"] = geometry["coordinates"][:2]
            if "id" not in row and "id" in feature:
                row["id"] = feature["id"]
            yield row
    else:
        with path.open(encoding=encoding, newline="") as file:
            sample = file.read(8192)
            file.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            yield from csv.DictReader(file, dialect=dialect)


def record(raw, source):
    data = {normalized(k): v for k, v in raw.items() if k is not None}
    def field(*names):
        for name in names:
            value = data.get(name)
            if value is not None and str(value).strip():
                return value
        return None
    identity = field("id", "id_unico", "id_camara", "id_poste")
    if identity is None:
        raise ValueError("Falta un identificador de origen; no se generan IDs ficticios")
    lat = float(field("latitude", "latitud", "lat", "coory"))
    lon = float(field("longitude", "longitud", "lon", "coorx"))
    if not math.isfinite(lat) or not math.isfinite(lon) or not (-85.05112878 <= lat <= 85.05112878 and -180 <= lon <= 180):
        raise ValueError("Coordenadas fuera del rango geográfico de Redis")
    return {"id": str(identity).strip(), "source": source,
            "direccion": field("direccion"), "esquina": field("esquina"), "colonia": field("colonia"),
            "alcaldia": field("alcaldia", "delegacion"), "poste": field("poste", "tipo_de_po", "tipo_de_poste"),
            "boton": field("boton"), "altavoz": field("altavoz"),
            "latitude": lat, "longitude": lon,
            "metadata": json.dumps(raw, ensure_ascii=False, default=str, allow_nan=False)}


def import_file(path, source, encoding="utf-8-sig"):
    counts = dict(procesados=0, insertados=0, actualizados=0, invalidos=0)
    with get_engine().begin() as c:
        c.execute(text("SELECT pg_advisory_xact_lock(:lock)"), {"lock": IMPORT_LOCK})
        postgis = c.execute(text("SELECT 1 FROM information_schema.columns WHERE table_schema='public' AND table_name='infraestructura_c5' AND column_name='location'")).scalar()
        for number, raw in enumerate(rows(path, encoding), start=2):
            counts["procesados"] += 1
            try:
                item = record(raw, source)
            except (ValueError, TypeError, OverflowError) as exc:
                counts["invalidos"] += 1
                print(f"Fila {number} inválida: {exc}", file=sys.stderr)
                continue
            exists = c.execute(text("SELECT source FROM infraestructura_c5 WHERE id=:id"), item).scalar()
            # An ID from another dataset must never silently overwrite its provenance.
            if exists is not None and exists != source:
                raise ValueError(f"ID {item['id']} ya pertenece a otra fuente; use un dataset o prefijo coherente")
            c.execute(text("""INSERT INTO infraestructura_c5(id,source,direccion,esquina,colonia,alcaldia,
              poste,boton,altavoz,latitude,longitude,metadata)
              VALUES(:id,:source,:direccion,:esquina,:colonia,:alcaldia,:poste,:boton,:altavoz,
              :latitude,:longitude,CAST(:metadata AS jsonb)) ON CONFLICT(id) DO UPDATE SET
              source=EXCLUDED.source,direccion=EXCLUDED.direccion,esquina=EXCLUDED.esquina,
              colonia=EXCLUDED.colonia,alcaldia=EXCLUDED.alcaldia,poste=EXCLUDED.poste,
              boton=EXCLUDED.boton,altavoz=EXCLUDED.altavoz,latitude=EXCLUDED.latitude,
              longitude=EXCLUDED.longitude,metadata=EXCLUDED.metadata,updated_at=clock_timestamp()"""), item)
            if postgis:
                c.execute(text("UPDATE infraestructura_c5 SET location=ST_SetSRID(ST_MakePoint(longitude,latitude),4326)::geography WHERE id=:id"), item)
            counts["actualizados" if exists is not None else "insertados"] += 1
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", nargs="?", type=Path)
    parser.add_argument("--source", help="Official dataset/resource URL, preserved as provenance")
    parser.add_argument("--encoding", default="utf-8-sig")
    parser.add_argument("--rebuild-redis", action="store_true")
    args = parser.parse_args()
    if args.file:
        if not args.source:
            parser.error("--source es obligatorio al importar un archivo")
        print(json.dumps(import_file(args.file, args.source, args.encoding), ensure_ascii=False))
    elif not args.rebuild_redis:
        parser.error("Proporciona archivo --source URL, o --rebuild-redis")
    try:
        print(json.dumps({"redis_indexados": rebuild_cache()}))
    except RedisError:
        print("PostgreSQL conservado. Redis pendiente; ejecuta --rebuild-redis al restablecerlo.", file=sys.stderr)
        raise SystemExit(2)


if __name__ == "__main__":
    raise SystemExit("Importador C5 archivado: no ejecutar. C5 usa el índice Redis existente; consulta docs/c5.md.")
